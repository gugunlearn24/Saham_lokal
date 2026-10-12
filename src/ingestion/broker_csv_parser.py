from __future__ import annotations

import csv
import json
import uuid
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from src.config.config_loader import load_brokers_config
from src.db.connection import get_connection
from src.ingestion.deduplication import compute_entry_hash, entry_hash_exists

DEFAULT_CONFIG_PATH = Path(__file__).resolve().parents[2] / "config" / "brokers.yaml"


def _parse_decimal(raw_value: Any, decimal_separator: str = ".", thousand_separator: str = ",") -> Decimal:
    if raw_value is None:
        raise ValueError("Value is missing.")

    text = str(raw_value).strip()
    if text == "":
        raise ValueError("Value is blank.")

    text = text.replace(thousand_separator, "")
    if decimal_separator != ".":
        text = text.replace(decimal_separator, ".")
    return Decimal(text)


def _normalize_trade_date(value: Any, date_format: str) -> str:
    text = str(value).strip()
    for fmt in (date_format, "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    raise ValueError(f"Unsupported trade date format: {value!r} (expected {date_format})")


def _normalize_action(raw_value: Any) -> str:
    text = str(raw_value).strip().upper()
    mapping = {
        "BUY": "BUY",
        "BELI": "BUY",
        "SELL": "SELL",
        "JUAL": "SELL",
    }
    if text not in mapping:
        normalized = text.replace(" ", "").replace("_", "")
        if normalized in mapping:
            return mapping[normalized]
        raise ValueError(f"Unsupported trade action: {raw_value!r}")
    return mapping[text]


def _normalize_ticker(value: Any) -> str:
    text = str(value).strip().upper()
    if not text:
        raise ValueError("Ticker is blank.")
    return text


def _resolve_broker_config(broker_id: str, config_path: str | Path = DEFAULT_CONFIG_PATH):
    brokers = load_brokers_config(config_path)
    for broker in brokers:
        if broker.get("broker_id") == broker_id:
            return broker
    raise ValueError(f"Unknown broker id '{broker_id}'. Expected one of: {[b.get('broker_id') for b in brokers]}")


def _read_rows_from_csv(
    file_path: str | Path,
    *,
    delimiter: str = ",",
    encoding: str = "utf-8",
) -> list[dict[str, Any]]:
    with Path(file_path).open("r", encoding=encoding, newline="") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def _read_rows_from_xlsx(file_path: str | Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]
    except ImportError as exc:
        raise RuntimeError("XLSX support requires openpyxl; install it with: pip install openpyxl") from exc

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook.active
    if sheet is None:
        workbook.close()
        raise ValueError("Workbook has no active worksheet.")
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    data_rows: list[dict[str, Any]] = []
    for row in rows[1:]:
        if not any(cell is not None and str(cell).strip() != "" for cell in row):
            continue
        item: dict[str, Any] = {}
        for idx, cell in enumerate(row):
            if idx < len(headers):
                item[headers[idx]] = cell
        data_rows.append(item)
    return data_rows


def load_trade_rows(file_path: str | Path, broker_id: str, config_path: str | Path = DEFAULT_CONFIG_PATH) -> list[dict[str, Any]]:
    broker = _resolve_broker_config(broker_id, config_path)
    path = Path(file_path)
    if path.suffix.lower() == ".csv":
        raw_rows = _read_rows_from_csv(
            path,
            delimiter=broker["delimiter"],
            encoding=broker["encoding"],
        )
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        raw_rows = _read_rows_from_xlsx(path)
    else:
        raise ValueError(f"Unsupported file format '{path.suffix}'. Expected CSV or XLSX.")

    mappings = broker["mappings"]
    conversions = broker.get("conversions", {})
    lot_multiplier = int(conversions.get("lot_multiplier", 1))
    decimal_separator = str(conversions.get("decimal_separator", "."))
    thousand_separator = str(conversions.get("thousand_separator", ","))

    normalized: list[dict[str, Any]] = []
    for index, row in enumerate(raw_rows, start=1):
        trade_date = _normalize_trade_date(row.get(mappings["trade_date"]), broker.get("date_format", "%Y-%m-%d"))
        ticker = _normalize_ticker(row.get(mappings["ticker"]))
        action = _normalize_action(row.get(mappings["action"]))

        shares_raw = row.get(mappings["shares"])
        shares_value = _parse_decimal(shares_raw, decimal_separator, thousand_separator)
        shares = int((shares_value * lot_multiplier).to_integral_value())

        price_raw = row.get(mappings["price_per_share"])
        price_per_share = float(_parse_decimal(price_raw, decimal_separator, thousand_separator))

        net_amount_raw = row.get(mappings["net_amount"])
        net_amount = int(_parse_decimal(net_amount_raw, decimal_separator, thousand_separator).to_integral_value())

        reference_id = str(row.get("reference_id") or row.get("Reference ID") or row.get("Order ID") or f"row-{index}")
        gross_amount = net_amount

        normalized.append(
            {
                "broker_id": broker["broker_id"],
                "trade_date": trade_date,
                "ticker": ticker,
                "action": action,
                "shares": shares,
                "price_per_share": price_per_share,
                "gross_amount": gross_amount,
                "net_cash_amount": net_amount,
                "reference_id": reference_id,
                "raw": row,
            }
        )
    return normalized


def ingest_broker_file(
    file_path: str | Path,
    broker_id: str,
    db_path: str | Path | None = None,
    config_path: str | Path = DEFAULT_CONFIG_PATH,
) -> dict[str, Any]:
    connection = get_connection(db_path) if db_path is not None else get_connection()
    rows = load_trade_rows(file_path, broker_id, config_path)

    inserted: list[str] = []
    skipped: list[dict[str, Any]] = []

    try:
        for row in rows:
            entry_hash = compute_entry_hash(
                row["broker_id"],
                row["trade_date"],
                row["ticker"],
                row["action"],
                row["shares"],
                row["gross_amount"],
                row["reference_id"],
            )

            if entry_hash_exists(connection, entry_hash):
                skipped.append({
                    "entry_hash": entry_hash,
                    "ticker": row["ticker"],
                    "trade_date": row["trade_date"],
                    "reference_id": row["reference_id"],
                    "reason": "duplicate",
                })
                continue

            staging_id = str(uuid.uuid4())
            payload = {
                "entry_hash": entry_hash,
                "broker_id": row["broker_id"],
                "trade_date": row["trade_date"],
                "ticker": row["ticker"],
                "action": row["action"],
                "shares": row["shares"],
                "price_per_share": row["price_per_share"],
                "gross_amount": row["gross_amount"],
                "reference_id": row["reference_id"],
                "raw": row["raw"],
            }
            connection.execute(
                """
                INSERT INTO reconciliation_staging (
                    staging_id, source_type, external_reference, trade_date, ticker, action,
                    shares, price_per_share, net_cash_amount, raw_payload, reconciliation_status
                ) VALUES (?, 'BROKER', ?, ?, ?, ?, ?, ?, ?, ?, 'PENDING')
                """,
                (
                    staging_id,
                    row["reference_id"],
                    row["trade_date"],
                    row["ticker"],
                    row["action"],
                    row["shares"],
                    row["price_per_share"],
                    row["net_cash_amount"],
                    json.dumps(payload, ensure_ascii=False, sort_keys=True),
                ),
            )
            inserted.append(staging_id)

        connection.commit()
        return {"inserted": inserted, "skipped": skipped}
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
