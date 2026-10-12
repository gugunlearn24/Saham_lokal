from __future__ import annotations

import csv
import json
import uuid
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from src.db.connection import get_connection


DATE_CANDIDATE_FORMATS = (
    "%Y-%m-%d",
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%m/%d/%Y",
    "%Y/%m/%d",
    "%d %b %Y",
    "%d %B %Y",
)


def _normalize_date(raw_value: Any) -> str:
    text = str(raw_value).strip()
    if not text:
        raise ValueError("Bank transaction date is required.")
    for fmt in DATE_CANDIDATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    raise ValueError(f"Unsupported bank transaction date format: {raw_value!r}")


def _parse_amount(raw_value: Any) -> int:
    if raw_value is None:
        return 0

    text = str(raw_value).strip()
    if text in {"", "-", "--"}:
        return 0

    is_negative = text.startswith("(") and text.endswith(")")
    if is_negative:
        text = "-" + text[1:-1]

    text = text.replace(" ", "").replace("_", "")
    text = text.replace("Rp", "").replace("IDR", "")
    text = text.replace(".", "") if "," in text and "." not in text else text
    if "," in text and "." in text:
        if text.rfind(".") > text.rfind(","):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(".", "").replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")

    if text.startswith("+"):
        text = text[1:]

    try:
        value = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Unable to parse bank amount: {raw_value!r}") from exc

    integer_value = int(value.quantize(Decimal("1")))
    return -integer_value if is_negative else integer_value


def _read_csv_rows(file_path: str | Path) -> list[dict[str, Any]]:
    with Path(file_path).open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _read_xlsx_rows(file_path: str | Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "XLSX support requires openpyxl; install it with: pip install openpyxl"
        ) from exc

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    data_rows: list[dict[str, Any]] = []
    for row in rows[1:]:
        if not any(cell is not None and str(cell).strip() != "" for cell in row):
            continue
        item: dict[str, Any] = {}
        for idx, cell in enumerate(row):
            if idx < len(headers):
                item[headers[idx]] = cell
        data_rows.append(item)
    return data_rows


def load_bank_rows(file_path: str | Path) -> list[dict[str, Any]]:
    path = Path(file_path)
    if path.suffix.lower() == ".csv":
        raw_rows = _read_csv_rows(path)
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        raw_rows = _read_xlsx_rows(path)
    else:
        raise ValueError(f"Unsupported bank file format '{path.suffix}'. Expected CSV or XLSX.")

    normalized_rows: list[dict[str, Any]] = []
    for index, row in enumerate(raw_rows, start=1):
        if not row:
            continue

        canonical_row = {
            str(key).strip().lower().replace(" ", "_").replace("-", "_"): value
            for key, value in row.items()
        }

        transaction_date = next(
            (
                _normalize_date(canonical_row.get(key))
                for key in (
                    "transaction_date",
                    "date",
                    "tanggal",
                    "trx_date",
                    "value_date",
                    "posting_date",
                )
                if key in canonical_row and canonical_row.get(key) not in (None, "")
            ),
            None,
        )
        if transaction_date is None:
            raise ValueError(f"Bank row #{index} is missing a recognizable date column.")

        debit_value = 0
        credit_value = 0
        for key in ("debit", "debit_amount", "nominal_debit", "debet", "withdrawal", "outflow"):
            if key in canonical_row and canonical_row.get(key) not in (None, ""):
                debit_value = _parse_amount(canonical_row.get(key))
                break

        for key in ("credit", "credit_amount", "nominal_credit", "kredit", "deposit", "inflow"):
            if key in canonical_row and canonical_row.get(key) not in (None, ""):
                credit_value = _parse_amount(canonical_row.get(key))
                break

        amount = credit_value - debit_value
        if amount == 0:
            continue

        description = next(
            (
                str(canonical_row.get(key)).strip()
                for key in ("description", "keterangan", "memo", "narration", "details")
                if key in canonical_row and canonical_row.get(key) not in (None, "")
            ),
            f"BANK-{index}",
        )

        reference_id = next(
            (
                str(canonical_row.get(key)).strip()
                for key in ("reference_id", "reference", "transaction_id", "trx_id", "id", "no_ref")
                if key in canonical_row and canonical_row.get(key) not in (None, "")
            ),
            description if description else f"bank-{index}",
        )

        normalized_rows.append(
            {
                "trade_date": transaction_date,
                "external_reference": reference_id,
                "description": description,
                "net_cash_amount": amount,
                "action": "DEPOSIT" if amount >= 0 else "WITHDRAW",
                "raw": canonical_row,
            }
        )

    return normalized_rows


def ingest_bank_statement(file_path: str | Path, db_path: str | Path | None = None) -> dict[str, Any]:
    connection = get_connection(db_path) if db_path is not None else get_connection()
    rows = load_bank_rows(file_path)
    inserted: list[str] = []

    try:
        for row in rows:
            staging_id = str(uuid.uuid4())
            payload = {
                "trade_date": row["trade_date"],
                "external_reference": row["external_reference"],
                "description": row["description"],
                "net_cash_amount": row["net_cash_amount"],
                "action": row["action"],
                "raw": row["raw"],
            }
            connection.execute(
                """
                INSERT INTO reconciliation_staging (
                    staging_id, source_type, external_reference, trade_date, ticker, action,
                    shares, price_per_share, net_cash_amount, raw_payload, reconciliation_status
                ) VALUES (?, 'BANK', ?, ?, NULL, ?, NULL, NULL, ?, ?, 'PENDING')
                """,
                (
                    staging_id,
                    row["external_reference"],
                    row["trade_date"],
                    row["action"],
                    row["net_cash_amount"],
                    json.dumps(payload, ensure_ascii=False, sort_keys=True),
                ),
            )
            inserted.append(staging_id)

        connection.commit()
        return {"inserted": inserted, "count": len(inserted)}
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()
